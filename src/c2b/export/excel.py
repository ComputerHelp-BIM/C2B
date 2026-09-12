"""Excel review workbook.

One sheet per element type plus Layers, Schedules, Tags and Diagnostics. The
workbook is the human review surface for step 1: filters on every sheet, frozen
headers, and the element id / DXF handles so anything can be traced back.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..schema import Project

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_SEV_FILL = {"ERROR": PatternFill("solid", fgColor="F8CBAD"), "WARNING": PatternFill("solid", fgColor="FFE699"), "INFO": PatternFill("solid", fgColor="DDEBF7")}


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: dict[int, float] | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([_cell(v) for v in r])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, start=1):
        best = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, best + 2), 60) if not widths or i not in widths else widths[i]
    return ws


def _cell(v):
    if isinstance(v, float):
        return round(v, 1)
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}={val}" for k, val in v.items())
    return v


def _outline_area(outline) -> float | None:
    if not outline or len(outline) < 3:
        return None
    pts = [(p.x, p.y) for p in outline]
    a = 0.0
    for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]):
        a += x1 * y2 - x2 * y1
    return abs(a) / 2 / 1e6


def _floor_name(project: Project, fid: str) -> str:
    return next((f.name for f in project.floors if f.id == fid), fid)


def write_workbook(project: Project, path: str | Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    s = project.summary
    d = project.drawing
    summary_rows = [
        ["File", d.file], ["DXF version", d.dxf_version], ["Units", f"{d.unit_name} (x{d.unit_scale_to_mm} to mm, {d.unit_source}, {d.unit_confidence})"],
        ["Generator", project.generator], ["Schema version", project.schema_version], ["Generated", project.generated_at], ["Profile", project.profile_name],
        [], ["Floors", s.floors], ["Grids", s.grids], ["Columns", s.columns], ["Beams", s.beams], ["Slabs", s.slabs], ["Footings", s.footings],
        ["Openings", s.openings], ["Walls", s.walls], ["Schedules", s.schedules], ["Unassigned tags", s.tags_unassigned],
        [], ["Errors", s.errors], ["Warnings", s.warnings], ["Infos", s.infos],
    ]
    _sheet(wb, "Summary", ["Item", "Value"], summary_rows)

    _sheet(wb, "Floors", ["Floor id", "Index", "Name", "Name source", "Origin X (drawing)", "Origin Y (drawing)", "Elevation (mm)", "Floor-to-floor (mm)",
                          "Default beam depth", "Default slab thk", "Grids", "Columns", "Beams", "Slabs", "Footings", "Openings", "Walls"],
           [[f.id, f.index, f.name, f.name_source, f.origin.x, f.origin.y, f.elevation_mm, f.floor_to_floor_mm, f.default_beam_depth_mm, f.default_slab_thickness_mm,
             f.counts.get("grids"), f.counts.get("columns"), f.counts.get("beams"), f.counts.get("slabs"), f.counts.get("footings"), f.counts.get("openings"), f.counts.get("walls")] for f in project.floors])

    _sheet(wb, "Layers", ["Layer", "Geometry role", "Text role", "Modifiers", "Confidence", "Source", "Layer size", "Entity counts"],
           [[l.layer, l.geometry_role, l.text_role, l.modifiers, l.confidence, l.source, l.layer_size, l.entity_counts] for l in project.layer_map])

    _sheet(wb, "Grids", ["Id", "Floor", "Label", "Axis", "Offset (mm)", "Start X", "Start Y", "End X", "End Y", "Angle", "Layer", "Handles"],
           [[g.id, _floor_name(project, g.floor_id), g.label, g.axis, g.offset_mm, g.start.x, g.start.y, g.end.x, g.end.y, g.angle_deg, g.source_layer, g.source_handles] for g in project.grids])

    _sheet(wb, "Columns", ["Id", "Floor", "Mark", "Grid ref", "Shape", "Center X", "Center Y", "Width (mm)", "Depth (mm)", "Dia (mm)", "Rotation", "Drawn W", "Drawn D",
                           "Size source", "Wall-like", "Modifier", "Tags", "Layer", "Source", "Confidence", "Handles"],
           [[c.id, _floor_name(project, c.floor_id), c.mark, c.grid_ref, c.shape, c.center.x, c.center.y, c.width_mm, c.depth_mm, c.diameter_mm, c.rotation_deg,
             c.drawn_width_mm, c.drawn_depth_mm, c.size_source, "yes" if c.wall_like else "", c.modifier, [t.text for t in c.tags], c.source_layer, c.source_kind, c.confidence, c.source_handles] for c in project.columns])

    _sheet(wb, "Beams", ["Id", "Floor", "Mark", "Start X", "Start Y", "End X", "End Y", "Length (mm)", "Width (mm)", "Depth (mm)", "Alt depth", "Drawn W", "Angle", "Inverted", "Sunk",
                         "Size source", "Depth source", "Tags", "Layer", "Source", "Edge parts", "Confidence", "Handles"],
           [[b.id, _floor_name(project, b.floor_id), b.mark, b.start.x, b.start.y, b.end.x, b.end.y, b.length_mm, b.width_mm, b.depth_mm, b.depth_alt_mm, b.drawn_width_mm, b.angle_deg,
             "yes" if b.inverted else "", b.sunk_mm, b.size_source, b.depth_source, [t.text for t in b.tags], b.source_layer, b.source_kind, b.n_edge_parts, b.confidence, b.source_handles] for b in project.beams])

    _sheet(wb, "Slabs", ["Id", "Floor", "Mark", "Thickness (mm)", "Thickness source", "Modifier", "Sunk", "Position X", "Position Y", "Has outline", "Area (m2)", "Tags", "Layer", "Source", "Handles"],
           [[sl.id, _floor_name(project, sl.floor_id), sl.mark, sl.thickness_mm, sl.thickness_source, sl.modifier, sl.sunk_mm, sl.position.x, sl.position.y, "yes" if sl.outline else "", _outline_area(sl.outline), [t.text for t in sl.tags], sl.source_layer, sl.source_kind, sl.source_handles] for sl in project.slabs])

    _sheet(wb, "Footings", ["Id", "Floor", "Mark", "Shape", "Center X", "Center Y", "Width (mm)", "Depth (mm)", "Thickness (mm)", "Fold (mm)", "Rotation", "Drawn W", "Drawn D", "Size source", "Modifier", "Tags", "Layer", "Handles"],
           [[f.id, _floor_name(project, f.floor_id), f.mark, f.shape, f.center.x, f.center.y, f.width_mm, f.depth_mm, f.thickness_mm, f.fold_mm, f.rotation_deg, f.drawn_width_mm, f.drawn_depth_mm, f.size_source, f.modifier, [t.text for t in f.tags], f.source_layer, f.source_handles] for f in project.footings])

    _sheet(wb, "Openings", ["Id", "Floor", "Label", "Center X", "Center Y", "Area (m2)", "Layer", "Handles"],
           [[o.id, _floor_name(project, o.floor_id), o.label, o.center.x, o.center.y, o.area_mm2 / 1e6, o.source_layer, o.source_handles] for o in project.openings])

    _sheet(wb, "Walls", ["Id", "Floor", "Mark", "Center X", "Center Y", "Length (mm)", "Thickness (mm)", "Rotation", "Structural", "Tags", "Layer", "Handles"],
           [[w.id, _floor_name(project, w.floor_id), w.mark, w.center.x, w.center.y, w.length_mm, w.thickness_mm, w.rotation_deg, "yes" if w.structural else "no", [t.text for t in w.tags], w.source_layer, w.source_handles] for w in project.walls])

    sched_rows = []
    for t in project.schedules:
        for r in t.rows:
            sched_rows.append([t.id, t.title, t.category, r.mark, r.values.get("width"), r.values.get("depth"), r.values.get("thickness"), r.values.get("diameter"), {k: v for k, v in r.values.items() if k not in ("width", "depth", "thickness", "diameter")}, t.source_layer])
    _sheet(wb, "Schedules", ["Schedule", "Title", "Category", "Mark", "Width", "Depth", "Thickness", "Dia", "Other", "Layer"], sched_rows)

    _sheet(wb, "Unassigned tags", ["Handle", "Floor", "Layer", "Role", "Text", "X", "Y", "Parsed"],
           [[t.handle, _floor_name(project, t.floor_id) if t.floor_id else "", t.layer, t.role, t.text, t.position.x, t.position.y,
             {k: v for k, v in t.parsed.items() if v not in (None, False, [], "") and k not in ("raw", "text")}] for t in project.tags_unassigned])

    ws = _sheet(wb, "Diagnostics", ["Severity", "Code", "Floor", "Layer", "Handle", "Element", "X", "Y", "Message"],
                [[d.severity, d.code, _floor_name(project, d.floor_id) if d.floor_id else "", d.layer, d.handle, d.element_id, d.location.x if d.location else None, d.location.y if d.location else None, d.message]
                 for d in sorted(project.diagnostics, key=lambda d: ({"ERROR": 0, "WARNING": 1, "INFO": 2}[d.severity], d.code))])
    for row in ws.iter_rows(min_row=2):
        fill = _SEV_FILL.get(row[0].value)
        if fill:
            row[0].fill = fill
    wb.save(str(path))
    return Path(path)
