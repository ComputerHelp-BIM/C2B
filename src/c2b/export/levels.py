"""Level schedule (Excel) template and reader.

Plans do not carry floor elevations. The extractor writes a template listing the
floors it found; the engineer fills in elevations; ``--levels`` feeds them back.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from ..schema import Project

HEADERS = ["floor_id", "floor_name", "order", "elevation_mm", "floor_to_floor_mm", "revit_level_name", "notes"]


def match_level_hint(floor_name: str, hints) -> tuple[float, str] | None:
    """Pick the level hint whose name shares the most words with the floor name."""
    words = {w for w in floor_name.upper().replace(".", " ").split() if w not in ("LEVEL", "LVL", "FLOOR", "AT", "LAYOUT", "PLAN", "-")}
    best, score = None, 0
    for h in hints:
        hw = {w for w in h.name.upper().replace(".", " ").split() if w not in ("LEVEL", "LVL", "FLOOR")}
        common = len(words & hw)
        if common > score:
            best, score = h, common
    return (best.elevation_mm, best.text) if best and score else None


def write_levels_template(project: Project, path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Levels"
    ws.append(HEADERS)
    for f in project.floors:
        hint = match_level_hint(f.name, project.level_hints) if f.elevation_mm is None else None
        elev = f.elevation_mm if f.elevation_mm is not None else (hint[0] if hint else None)
        note = "fill elevation_mm (top of structural slab) in mm" if elev is None else (f"from client text '{hint[1]}', please confirm" if hint else "")
        ws.append([f.id, f.name, f.index, elev, f.floor_to_floor_mm, None, note])
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFG", (10, 36, 8, 16, 18, 24, 50)):
        ws.column_dimensions[col].width = width
    st = wb.create_sheet("Settings")
    st.append(["key", "value", "notes"])
    st.append(["level_reference", "SSL", "SSL = top of structural slab (default), FFL = finished floor; this sheet wins over the template spec"])
    st.column_dimensions["A"].width = 18
    st.column_dimensions["B"].width = 12
    st.column_dimensions["C"].width = 80
    if project.level_hints:
        hs = wb.create_sheet("Level hints")
        hs.append(["name", "elevation_mm", "client text", "handle", "layer", "floor"])
        for h in sorted(project.level_hints, key=lambda h: h.elevation_mm):
            hs.append([h.name, h.elevation_mm, h.text, h.handle, h.layer, h.floor_id])
        for col, width in zip("ABCDEF", (30, 14, 60, 10, 24, 8)):
            hs.column_dimensions[col].width = width
    wb.save(str(path))
    return Path(path)


def apply_levels(project: Project, path: str | Path) -> list[str]:
    """Read a filled level schedule into the project floors. Returns a list of problems."""
    problems: list[str] = []
    wb = load_workbook(str(path), data_only=True)
    ws = wb["Levels"] if "Levels" in wb.sheetnames else wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    if "floor_id" not in idx:
        return ["level schedule has no 'floor_id' column"]
    by_id = {f.id: f for f in project.floors}
    by_name = {f.name.strip().lower(): f for f in project.floors}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        fid = str(row[idx["floor_id"]]).strip() if row[idx["floor_id"]] is not None else ""
        floor = by_id.get(fid) or by_name.get(str(row[idx.get("floor_name", 0)] or "").strip().lower())
        if floor is None:
            problems.append(f"level row '{fid}' does not match any floor in the drawing")
            continue
        for key, attr in (("elevation_mm", "elevation_mm"), ("floor_to_floor_mm", "floor_to_floor_mm")):
            if key in idx and row[idx[key]] is not None:
                try:
                    setattr(floor, attr, float(row[idx[key]]))
                except (TypeError, ValueError):
                    problems.append(f"{fid}: {key} '{row[idx[key]]}' is not a number")
        if "order" in idx and row[idx["order"]] is not None:
            try:
                floor.index = int(row[idx["order"]])
            except (TypeError, ValueError):
                problems.append(f"{fid}: order '{row[idx['order']]}' is not an integer")
    project.floors.sort(key=lambda f: f.index)
    return problems


def read_levels(path: str | Path) -> list:
    """Read the level schedule as rows (a plan floor may own several levels, e.g. a typical floor)."""
    from ..normalize.pipeline import LevelRow

    wb = load_workbook(str(path), data_only=True)
    ws = wb["Levels"] if "Levels" in wb.sheetnames else wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue

        def val(key, row=row):
            i = idx.get(key)
            return row[i] if i is not None and i < len(row) else None

        def num(key):
            v = val(key)
            try:
                return float(v) if v is not None and str(v).strip() != "" else None
            except (TypeError, ValueError):
                return None

        fid = str(val("floor_id")).strip() if val("floor_id") not in (None, "") else None
        rows.append(LevelRow(fid, str(val("floor_name") or "").strip() or None, int(num("order") or 0) if num("order") is not None else None,
                             num("elevation_mm"), num("floor_to_floor_mm"), str(val("revit_level_name") or "").strip() or None))
    return rows


def level_rows_for_editing(path: str | Path) -> list[dict]:
    """The level sheet as plain rows, for an editor that is not Excel.

    A drafter should not have to leave the window, find a workbook, type into it and come back:
    the elevations are the one thing the drawing cannot supply, and asking for them in the app
    is the difference between one run and three.
    """
    wb = load_workbook(str(path), data_only=True)
    ws = wb["Levels"] if "Levels" in wb.sheetnames else wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    rows = []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v is None for v in row):
            continue

        def cell(key, row=row):
            i = idx.get(key)
            return row[i] if i is not None and i < len(row) else None

        elevation = cell("elevation_mm")
        rows.append({"excel_row": n,
                     "floor_id": str(cell("floor_id") or "").strip(),
                     "floor_name": str(cell("floor_name") or "").strip(),
                     "revit_level_name": str(cell("revit_level_name") or "").strip(),
                     "elevation_mm": float(elevation) if isinstance(elevation, (int, float)) else None,
                     "notes": str(cell("notes") or "").strip()})
    return rows


def write_level_elevations(path: str | Path, elevations: dict[int, float | None]) -> int:
    """Write elevations back by worksheet row, keeping everything else in the workbook."""
    wb = load_workbook(str(path))
    ws = wb["Levels"] if "Levels" in wb.sheetnames else wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    if "elevation_mm" not in header:
        raise ValueError("this level workbook has no 'elevation_mm' column")
    column = header.index("elevation_mm") + 1
    written = 0
    for excel_row, value in elevations.items():
        ws.cell(row=excel_row, column=column).value = value
        written += 1
    wb.save(str(path))
    return written


def read_level_settings(path: str | Path) -> dict[str, str]:
    """Key/value pairs from the workbook's Settings sheet (e.g. level_reference)."""
    wb = load_workbook(str(path), data_only=True)
    if "Settings" not in wb.sheetnames:
        return {}
    out: dict[str, str] = {}
    for row in wb["Settings"].iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None and len(row) > 1 and row[1] is not None:
            out[str(row[0]).strip()] = str(row[1]).strip()
    return out


def write_levels_from_storeys(schedule, path: str | Path, project: Project | None = None,
                              level_reference: str = "SSL") -> Path:
    """The level workbook, written from the storey schedule.

    The storey window owns the building's levels; this workbook is what it produces. It keeps
    the columns :func:`read_levels` has always read, so a scripted run pointing ``--levels`` at
    it still works and Revit still gets its levels from one place -- but it is an *output* now,
    and the banner on the sheet says so, because a drafter who types into it would lose the
    edit on the next run. To work in Excel instead, delete the ``.storeys.json`` beside it.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Levels"
    ws.append(HEADERS)
    rows = schedule.rows()
    for r, s in zip(rows, schedule.storeys):
        note = "; ".join(x for x in (s.source_words, s.note) if x)
        ws.append([s.plan_floor_id, s.name, r["n"], r["elevation_mm"],
                   None if r["is_top"] else schedule.storeys[r["n"] + 1].height_mm, s.name, note])
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFG", (10, 36, 8, 16, 18, 24, 50)):
        ws.column_dimensions[col].width = width

    st = wb.create_sheet("Settings")
    st.append(["key", "value", "notes"])
    st.append(["level_reference", level_reference,
               "SSL = top of structural slab (default), FFL = finished floor; this sheet wins over the template spec"])
    st.append(["written_by", "the C2B storey window",
               "This workbook is written from the storey schedule on every run. Edit the storeys in "
               "C2B, not here: anything typed into this file is replaced the next time you press Run."])
    st.append(["base_elevation_mm", schedule.base_elevation_mm, "elevation of the lowest storey; every other is this plus the heights above it"])
    st.append(["total_height_mm", schedule.total_height_mm(), "lowest storey to highest"])
    st.column_dimensions["A"].width = 20
    st.column_dimensions["B"].width = 26
    st.column_dimensions["C"].width = 100

    if project is not None and project.level_hints:
        hs = wb.create_sheet("Level hints")
        hs.append(["name", "elevation_mm", "client text", "handle", "layer", "floor"])
        for h in sorted(project.level_hints, key=lambda h: h.elevation_mm):
            hs.append([h.name, h.elevation_mm, h.text, h.handle, h.layer, h.floor_id])
        for col, width in zip("ABCDEF", (30, 14, 60, 10, 24, 8)):
            hs.column_dimensions[col].width = width
    wb.save(str(path))
    return Path(path)
