"""levels.xlsx is written from the storey schedule and still reads back the same way."""
from __future__ import annotations

from c2b.export.levels import read_level_settings, read_levels, write_levels_from_storeys
from c2b.storeys import StoreySchedule


def stack(*heights, base=0.0, names=None):
    s = StoreySchedule(base_elevation_mm=base)
    for i, h in enumerate(heights):
        s.add(name=(names[i] if names else f"L{i}"), height_mm=h, plan_floor_id=f"F{i:02d}", source="cad")
    return s


def test_the_workbook_reads_back_as_the_stack_it_was_written_from(tmp_path):
    s = stack(0, 3000, 3500, base=-2500, names=["FOUNDATION", "GF", "1F"])
    path = write_levels_from_storeys(s, tmp_path / "x.levels.xlsx")
    rows = read_levels(path)
    assert [r.elevation for r in rows] == [-2500, 500, 4000]
    assert [r.revit_name for r in rows] == ["FOUNDATION", "GF", "1F"]
    assert [r.floor_id for r in rows] == ["F00", "F01", "F02"]
    assert [r.order for r in rows] == [0, 1, 2]


def test_floor_to_floor_is_the_rise_to_the_storey_above(tmp_path):
    s = stack(0, 3000, 4000)
    rows = read_levels(write_levels_from_storeys(s, tmp_path / "x.xlsx"))
    assert [r.f2f for r in rows] == [3000, 4000, None]


def test_typical_floors_all_carry_the_one_plan_id(tmp_path):
    s = stack(0, 3000, names=["GF", "1F"])
    s.repeat(1, times=3)
    rows = read_levels(write_levels_from_storeys(s, tmp_path / "x.xlsx"))
    assert [r.floor_id for r in rows] == ["F00", "F01", "F01", "F01", "F01"]
    assert [r.elevation for r in rows] == [0, 3000, 6000, 9000, 12000]


def test_the_workbook_says_it_is_written_not_filled_in(tmp_path):
    s = stack(0, 3000)
    settings = read_level_settings(write_levels_from_storeys(s, tmp_path / "x.xlsx"))
    assert settings["level_reference"] == "SSL"
    assert "storey window" in settings["written_by"]
    assert float(settings["base_elevation_mm"]) == 0.0


def test_the_level_reference_asked_for_is_the_one_written(tmp_path):
    settings = read_level_settings(write_levels_from_storeys(stack(0), tmp_path / "x.xlsx", level_reference="FFL"))
    assert settings["level_reference"] == "FFL"


def test_each_row_says_where_its_storey_came_from(tmp_path):
    s = stack(0, 3000, names=["GF", "1F"])
    s.repeat(1, times=1)
    rows = read_levels(write_levels_from_storeys(s, tmp_path / "x.xlsx"))
    assert len(rows) == 3
    from openpyxl import load_workbook
    ws = load_workbook(str(tmp_path / "x.xlsx"))["Levels"]
    notes = [ws.cell(row=r, column=7).value for r in (2, 3, 4)]
    assert "floor plan in the drawing" in notes[0]
    assert "repeat" in notes[2]


def test_an_empty_schedule_writes_a_workbook_with_only_headers(tmp_path):
    path = write_levels_from_storeys(StoreySchedule(), tmp_path / "x.xlsx")
    assert read_levels(path) == []
