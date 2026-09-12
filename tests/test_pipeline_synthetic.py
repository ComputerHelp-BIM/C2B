"""End-to-end test on the synthetic two-floor drawing built in conftest."""
from c2b.export.excel import write_workbook
from c2b.export.jsonout import read_json, write_json
from c2b.export.levels import apply_levels, write_levels_template
from c2b.export.report import write_report
from c2b.export.review_dxf import write_review_dxf


def test_floors_detected(synthetic_result):
    p = synthetic_result.project
    assert [f.name for f in p.floors] == ["FOUNDATION LEVEL", "GROUND FLOOR LEVEL"]
    assert p.floors[0].origin.x == 0 and p.floors[1].origin.x == 40000
    assert p.drawing.unit_name == "millimetres"


def test_grids(synthetic_result):
    p = synthetic_result.project
    g2 = [g for g in p.grids if g.floor_id == "L02"]
    assert len(g2) == 7
    assert sorted(g.label for g in g2 if g.axis == "X") == ["1", "2", "3", "4"]
    assert sorted(g.label for g in g2 if g.axis == "Y") == ["A", "B", "C"]
    assert any(g.label == "2" and abs(g.offset_mm - 6000) < 1 for g in g2)


def test_columns(synthetic_result):
    p = synthetic_result.project
    cols = [c for c in p.columns if c.floor_id == "L02"]
    assert len(cols) == 12
    circ = [c for c in cols if c.shape == "circle"]
    assert len(circ) == 1 and circ[0].diameter_mm == 600 and circ[0].size_source == "tag"
    from_lines = [c for c in cols if c.source_kind == "lines"]
    assert len(from_lines) == 1 and from_lines[0].mark == "C1" and from_lines[0].size_source == "schedule"
    assert from_lines[0].width_mm == 300 and from_lines[0].depth_mm == 450
    tagged = [c for c in cols if c.size_source == "tag" and c.shape == "rect"]
    assert len(tagged) == 10 and all((c.width_mm, c.depth_mm) == (300, 450) for c in tagged)
    assert all(c.grid_ref for c in cols)
    assert {c.grid_ref for c in cols} >= {"1/A", "4/C"}


def test_beams(synthetic_result):
    p = synthetic_result.project
    beams = [b for b in p.beams if b.floor_id == "L02"]
    assert len(beams) == 8
    horiz = [b for b in beams if abs(b.angle_deg) < 1]
    vert = [b for b in beams if abs(b.angle_deg - 90) < 1]
    assert len(horiz) == 4 and len(vert) == 4
    assert all(abs(b.drawn_width_mm - 230) < 1 for b in beams)
    tagged = [b for b in horiz if b.mark is None and b.depth_mm == 450 and b.depth_source == "tag"]
    assert len(tagged) == 3
    default = [b for b in horiz if b.depth_source == "default"]
    assert len(default) == 1 and default[0].depth_mm == 450
    assert all(b.mark == "MB" and b.depth_mm == 600 for b in vert), [(b.mark, b.depth_mm, b.depth_source) for b in vert]
    assert not [b for b in p.beams if b.floor_id == "L01"]


def test_footings_slabs_openings(synthetic_result):
    p = synthetic_result.project
    ftg = [f for f in p.footings if f.floor_id == "L01"]
    assert len(ftg) == 12 and all(f.mark == "F1" and f.thickness_mm == 500 for f in ftg)
    slabs = [s for s in p.slabs if s.floor_id == "L02"]
    assert len(slabs) == 6 and all(s.thickness_mm == 150 for s in slabs)
    ops = [o for o in p.openings if o.floor_id == "L02"]
    assert len(ops) == 1 and ops[0].label == "LIFT"


def test_schedule_and_diagnostics(synthetic_result):
    p = synthetic_result.project
    assert len(p.schedules) == 1 and p.schedules[0].category == "column"
    assert p.summary.errors == 0
    codes = {d.code for d in p.diagnostics}
    assert "SLAB_NONE_ON_FLOOR" in codes            # foundation floor has no slab tags
    assert "DEFAULT_APPLIED" in codes
    assert not [d for d in p.diagnostics if d.code == "TAG_UNASSIGNED"], [d.message for d in p.diagnostics if d.code == "TAG_UNASSIGNED"]


def test_exports_roundtrip(synthetic_result, tmp_path):
    p = synthetic_result.project
    j = write_json(p, tmp_path / "x.c2b.json")
    again = read_json(j)
    assert again.summary.columns == p.summary.columns and again.schema_version == p.schema_version
    write_workbook(p, tmp_path / "x.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "x.xlsx")
    assert {"Summary", "Floors", "Layers", "Columns", "Beams", "Diagnostics"} <= set(wb.sheetnames)
    assert wb["Columns"].max_row == p.summary.columns + 1
    write_report(p, tmp_path / "x.md")
    assert "## Diagnostics by code" in (tmp_path / "x.md").read_text()
    dxf = write_review_dxf(p, tmp_path / "x.review.dxf")
    import ezdxf
    doc = ezdxf.readfile(str(dxf))
    assert len(doc.modelspace().query('LWPOLYLINE[layer=="C2B-COLUMN"]')) == p.summary.columns


def test_levels_template_and_apply(synthetic_result, tmp_path):
    p = synthetic_result.project
    t = write_levels_template(p, tmp_path / "levels.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(t)
    ws = wb["Levels"]
    ws["D2"], ws["D3"], ws["E2"] = -1500, 0, 1500
    wb.save(t)
    problems = apply_levels(p, t)
    assert problems == []
    assert p.floors[0].elevation_mm == -1500 and p.floors[1].elevation_mm == 0 and p.floors[0].floor_to_floor_mm == 1500
